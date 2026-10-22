"""Prevod mesačných Excelov so smenovkou na dáta pre web.

Spustenie:
    python sync.py          sleduje priečinok s Excelmi a po každom uložení aktualizuje web
    python sync.py --once   spracuje všetko raz a skončí
"""
import datetime
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

# ------------------------- NASTAVENIA -------------------------
EXCEL_FOLDER = Path(r"C:\Users\Gastro\Desktop\PLANOVANIE A ODMENY NEW")
PUSH = True                               # True = po zmene sám nahrá web na GitHub
PUSH_EVERY_MINUTES = 0                    # 0 = odosiela zmeny ihneď po uložení
POLL_SECONDS = 20                         # ako často kontroluje priečinok
# --------------------------------------------------------------

BASE = Path(__file__).resolve().parent
SITE = BASE / "docs"
DATA = SITE / "data"
FILE_RE = re.compile(r"^Master_Smenovka_(\d{4})_(\d{2})\.xlsx$", re.I)
MONTHS = ["Január", "Február", "Marec", "Apríl", "Máj", "Jún", "Júl",
          "August", "September", "Október", "November", "December"]


def plain(v):
    """Text bez diakritiky, veľkými písmenami (na porovnávanie)."""
    s = unicodedata.normalize("NFKD", str(v or ""))
    return "".join(c for c in s if not unicodedata.combining(c)).strip().upper()


def find_sheet(wb, *words):
    for ws in wb.worksheets:
        if all(w in plain(ws.title) for w in words):
            return ws
    raise ValueError("V súbore chýba list: " + " ".join(words))


def read_schedule(ws, year, month):
    cols = {}
    for c in range(3, ws.max_column + 1):
        v = ws.cell(1, c).value
        if isinstance(v, datetime.datetime) and v.year == year and v.month == month:
            cols[v.day] = c
    days = sorted(cols)
    shifts = []
    for r in range(3, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name or plain(ws.cell(r, 2).value) != "SMENA":
            continue
        s = {}
        for d in days:
            v = ws.cell(r, cols[d]).value
            if v not in (None, ""):
                s[str(d)] = str(v).strip()
        # pracovníci, ktorí nemajú v mesiaci žiadnu smenu (len "x" alebo prázdno), sa vynechajú
        if any(not code.upper().startswith("X") for code in s.values()):
            shifts.append({"n": str(name).strip(), "s": s})
    return days, shifts


def read_rewards(ws, key):
    """Nájde blok odmien podľa nadpisu v stĺpci B (napr. VNUTORNA OBSLUHA)."""
    marker = next((r for r in range(1, ws.max_row + 1)
                   if plain(ws.cell(r, 2).value).startswith(key)), None)
    if marker is None:
        return []
    cat_row = marker + 3
    cats = [(str(ws.cell(cat_row, c).value).strip(), c)
            for c in range(8, ws.max_column + 1) if ws.cell(cat_row, c).value]
    out = []
    r = marker + 5
    while r <= ws.max_row and ws.cell(r, 2).value:
        name = str(ws.cell(r, 2).value).strip()
        num = lambda row, col: round(float(ws.cell(row, col).value or 0), 4)
        total, fulfil = num(r, 5), num(r, 4)
        if total or fulfil:                       # brigádnici s nulami sa nezobrazia
            rank = ws.cell(r, 3).value
            out.append({
                "n": name,
                "rank": rank if isinstance(rank, (int, float)) else None,
                "tot": round(total, 2), "pl": fulfil,
                "c": [[cn.title() if cn.isupper() else cn,
                       num(r, c), num(r, c + 1), num(r, c + 2),
                       round(num(r, c + 3), 2), round(num(r, c + 4), 2)]
                      for cn, c in cats],
            })
        r += 1
    out.sort(key=lambda x: (x["rank"] is None, x["rank"] or 0))
    return out


def base_cat(name):
    """Názov kategórie bez zátvorky na konci ('Cigarety (noc)' -> 'Cigarety')."""
    return re.sub(r"\s*\(.*?\)\s*$", "", str(name)).strip()


def read_plans(ws):
    """Denný plán tovaru podľa typu smeny (D-VO / D / N) z listu PLANY.

    List PLANY má sekcie:
        --- 1. VONKAJŠIA SMENA (D-VO) ---   stĺpec A = tovar, B = merná jednotka, C = denný plán
        --- 2. VNÚTORNÁ DENNÁ SMENA (D) ---  stĺpec D = váha kategórie (v %)
        --- 3. VNÚTORNÁ NOČNÁ SMENA (N) ---  váha sa berie z dennej sekcie
    Vrací {"D": {"CIGARETY": {"p": 150, "u": "ks", "w": 10}}, "N": {...}, "D-VO": {...}}
    """
    plans = {}
    kind = None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a in (None, ""):
            continue
        s = str(a).strip()
        if s.startswith("---"):                      # nadpis sekcie
            up = plain(s)
            if "D-VO" in up:
                kind = "D-VO"
            elif "NOC" in up:
                kind = "N"
            elif "DENN" in up:
                kind = "D"
            else:
                kind = None
            continue
        if not kind:
            continue
        val = ws.cell(r, 3).value
        if not isinstance(val, (int, float)):
            continue
        cat = base_cat(s)
        if not cat:
            continue
        entry = {
            "p": round(float(val), 4),
            "u": str(ws.cell(r, 2).value or "").strip(),
        }
        weight = ws.cell(r, 4).value            # stĺpec D = Váha kategórie (v %)
        if isinstance(weight, (int, float)):
            entry["w"] = round(float(weight), 4)
        plans.setdefault(kind, {})[plain(cat)] = entry

    # Nočná sekcia nemá v PLANY vyplnenú váhu – platí rovnaká váha ako pri dennej
    # (jeden tovar má v Exceli jednu váhu, tá istá pre D aj N).
    for name, entry in plans.get("N", {}).items():
        if "w" not in entry and "w" in plans.get("D", {}).get(name, {}):
            entry["w"] = plans["D"][name]["w"]
    return plans


def read_month(path, year, month):
    with tempfile.TemporaryDirectory() as tmp:    # kópia, aby nevadil otvorený Excel
        copy = Path(tmp) / "x.xlsx"
        shutil.copy2(path, copy)
        wb = load_workbook(copy, data_only=True)
    days, shifts = read_schedule(find_sheet(wb, "MASTER"), year, month)
    su = find_sheet(wb, "SUMAR")
    try:
        plans = read_plans(find_sheet(wb, "PLANY"))
    except Exception:
        plans = {}
    return {
        "id": f"{year}-{month:02d}", "label": f"{MONTHS[month - 1]} {year}",
        "y": year, "m": month - 1, "days": days, "shifts": shifts,
        "vno": read_rewards(su, "VNUTORNA"), "vo": read_rewards(su, "VONKAJSIA"),
        "plans": plans,
    }


def digest(d):
    return hashlib.md5(json.dumps(d, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


CHANGE_HOURS = 48                         # ako dlho sa zmena na webe zvýrazňuje


def merge_changes(old_full, data, today):
    """Zmenené smeny (dnes a v budúcnosti) + pôvodná hodnota; staré záznamy po 48 h zmiznú."""
    now = datetime.datetime.now().astimezone()
    limit = datetime.timedelta(hours=CHANGE_HOURS)
    ch = {}
    for name, days in (old_full.get("changes") or {}).items():
        for day, info in days.items():
            try:
                fresh = now - datetime.datetime.fromisoformat(info["at"]) < limit
                future = datetime.date(data["y"], data["m"] + 1, int(day)) >= today
            except Exception:
                continue
            if fresh and future:
                ch.setdefault(name, {})[day] = info
    old_map = {e["n"]: e["s"] for e in old_full.get("shifts", [])}
    if old_map:                                   # nový mesiac sa neoznačuje
        for e in data["shifts"]:
            if e["n"] not in old_map:             # nový zamestnanec sa neoznačuje
                continue
            for d in data["days"]:
                new, prev = e["s"].get(str(d), ""), old_map[e["n"]].get(str(d), "")
                if new != prev and datetime.date(data["y"], data["m"] + 1, d) >= today:
                    ch.setdefault(e["n"], {})[str(d)] = {"was": prev or "-", "at": now.isoformat(timespec="seconds")}
    return ch


def process_all():
    DATA.mkdir(parents=True, exist_ok=True)
    changed, months = False, []

    # Aktuálny rok a mesiac: 1 mesiac dozadu až 2 mesiace dopredu (-1, 0, +1, +2)
    today = datetime.date.today()
    curr_idx = today.year * 12 + (today.month - 1)
    valid_range = {curr_idx - 1, curr_idx, curr_idx + 1, curr_idx + 2}

    for f in sorted(EXCEL_FOLDER.glob("Master_Smenovka_*.xlsx")):
        m = FILE_RE.match(f.name)
        if not m:
            continue
        year, month = int(m.group(1)), int(m.group(2))
        f_idx = year * 12 + (month - 1)

        # Vyfiltruj iba povolené mesiace (-1 až +2)
        if f_idx not in valid_range:
            continue

        data = read_month(f, year, month)
        months.append({"id": data["id"], "label": data["label"]})
        target = DATA / f"{data['id']}.json"
        old_full = json.loads(target.read_text("utf-8")) if target.exists() else {}
        old = dict(old_full)
        old.pop("updated", None)
        old.pop("changes", None)
        if digest(old) != digest(data):
            data["changes"] = merge_changes(old_full, data, today)
            data["updated"] = datetime.datetime.now().astimezone().isoformat(timespec="seconds")
            target.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), "utf-8")
            print(time.strftime("%H:%M:%S"), "aktualizované:", f.name)
            changed = True

    index = DATA / "index.json"
    new_index = json.dumps({"months": months}, ensure_ascii=False)
    if not index.exists() or index.read_text("utf-8") != new_index:
        index.write_text(new_index, "utf-8")
        changed = True
    return changed


def push():
    for cmd in (["git", "add", "-A"],
                ["git", "commit", "-m", "aktualizácia smenovky"],
                ["git", "push"]):
        r = subprocess.run(cmd, cwd=BASE, capture_output=True, text=True)
        if r.returncode and "nothing to commit" not in (r.stdout + r.stderr):
            print("Git zlyhal:", (r.stderr or r.stdout).strip())
            return False
    print(time.strftime("%H:%M:%S"), "web nahraný, na adrese bude za minútu-dve")
    return True


def signature():
    return tuple((p.name, p.stat().st_mtime) for p in sorted(EXCEL_FOLDER.glob("Master_Smenovka_*.xlsx")))


def main():
    once = "--once" in sys.argv
    last, pending, last_push = None, False, 0.0
    print("Sledujem priečinok:", EXCEL_FOLDER)
    while True:
        try:
            sig = signature()
            settled = all(time.time() - mt > 5 for _, mt in sig)   # počká na dokončenie ukladania v Exceli
            if sig != last and settled:
                if process_all():
                    pending = True
                last = sig
            due = time.time() - last_push >= PUSH_EVERY_MINUTES * 60
            if pending and PUSH and (due or once):
                if push():
                    pending, last_push = False, time.time()
        except Exception as e:
            print(time.strftime("%H:%M:%S"), "čakám:", e)
        if once:
            break
        time.sleep(POLL_SECONDS)


if __name__ == "__main__":
    main()
